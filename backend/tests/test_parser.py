"""解析层单测：最小 docx、最小文本型 pdf、无文本层 pdf，覆盖三条异常路径。"""

import pytest
from docx import Document
from fpdf import FPDF

from app.errors import EmptyFileError, NoTextLayerError, UnsupportedTypeError
from app.services.parser import parse_file


def test_parse_docx(tmp_path):
    p = tmp_path / "a.docx"
    doc = Document()
    doc.add_paragraph("实验目的")
    doc.add_paragraph(
        "掌握单链表的存储结构与基本操作，理解指针在动态内存分配中的作用，"
        "学会用单链表解决顺序存储不便插入删除的问题，体会链式存储与顺序存储的差异"
    )
    doc.save(str(p))
    r = parse_file(str(p))
    assert r["source_type"] == "docx"
    assert "实验目的" in r["text"]
    assert r["char_count"] > 0


def test_parse_pdf_text(tmp_path):
    p = tmp_path / "a.pdf"
    pdf = FPDF()
    pdf.add_page()
    pdf.set_font("helvetica", size=12)
    pdf.multi_cell(
        0,
        10,
        "Experiment purpose: master the storage structure and basic operations "
        "of a singly linked list. Understand the role of pointers in dynamic "
        "memory allocation. Experiment principle: a singly linked list consists "
        "of nodes, each node has a data field and a next pointer.",
    )
    pdf.output(str(p))
    r = parse_file(str(p))
    assert r["source_type"] == "pdf"
    assert r["char_count"] >= 50


def test_parse_pdf_no_text(tmp_path):
    p = tmp_path / "scan.pdf"
    pdf = FPDF()
    pdf.add_page()  # 空白页，无文本层
    pdf.output(str(p))
    with pytest.raises(NoTextLayerError) as e:
        parse_file(str(p))
    assert e.value.code == 4003


def test_parse_unsupported_ext(tmp_path):
    p = tmp_path / "a.txt"
    p.write_text("hello", encoding="utf-8")
    with pytest.raises(UnsupportedTypeError) as e:
        parse_file(str(p))
    assert e.value.code == 4001


def test_parse_empty(tmp_path):
    p = tmp_path / "a.docx"
    p.write_bytes(b"")
    with pytest.raises(EmptyFileError) as e:
        parse_file(str(p))
    assert e.value.code == 4004


def test_parse_corrupt_docx(tmp_path):
    p = tmp_path / "bad.docx"
    p.write_bytes(b"this is not a real docx zip file")
    with pytest.raises(NoTextLayerError) as e:
        parse_file(str(p))
    assert e.value.code == 4003
    assert "docx" in e.value.message and "PDF" not in e.value.message


def test_parse_docx_no_text(tmp_path):
    p = tmp_path / "empty.docx"
    Document().save(str(p))  # 空文档，抽不出文字
    with pytest.raises(NoTextLayerError) as e:
        parse_file(str(p))
    assert e.value.code == 4003
    assert "docx" in e.value.message and "PDF" not in e.value.message


def test_paragraphs_coordinates_self_consistent(tmp_path):
    """段落坐标自洽：逐段 text[start:end] 相等、相邻段间隔为分隔符、末段 end == len(text)。"""
    p = tmp_path / "a.docx"
    doc = Document()
    doc.add_paragraph("第一段：掌握单链表的存储结构与基本操作")
    doc.add_paragraph("第二段：理解指针在动态内存分配中的作用")
    doc.add_paragraph("第三段：学会用单链表解决顺序存储插入删除不便的问题")
    doc.save(str(p))
    r = parse_file(str(p))
    text = r["text"]
    paras = r["paragraphs"]

    assert len(paras) == 3
    for para in paras:
        assert text[para["start"]:para["end"]] == para["text"]
    for i in range(len(paras) - 1):
        assert paras[i]["end"] + 1 == paras[i + 1]["start"]  # 中间的正是分隔符
    assert paras[-1]["end"] == len(text)
